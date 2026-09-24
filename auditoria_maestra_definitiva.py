"""
AUDITORÍA MAESTRA DEFINITIVA: 04-SEP-2026 (Versión 12:21 PM)
Verificación exhaustiva de hoja VENTAS, cruces con LISTADO y PROYECCIÓN
"""
import openpyxl, sys, re, calendar
from datetime import datetime, date

sys.stdout.reconfigure(encoding='utf-8')

FILE = 'Base de datos - La Ceiba ultima versión 2208.xlsx'
print(f"Cargando {FILE}...")
wb = openpyxl.load_workbook(FILE, data_only=True)

ws_l = wb['LISTADO']
ws_v = wb['VENTAS']
p_name = [s for s in wb.sheetnames if 'PROY' in s.upper()][0]
ws_p = wb[p_name]

def to_num(val):
    if val is None: return 0.0
    try:
        if isinstance(val, (int, float)): return float(val)
        s = str(val).replace('$','').replace(',','').strip()
        if s.startswith('#') or s in ('','None','nan'): return 0.0
        return float(s)
    except: return 0.0

def clean_str(val):
    if val is None: return None
    s = str(val).strip()
    return s if s and s.lower() not in ('none','nan','') else None

def norm(val):
    if not val: return ''
    s = str(val).strip().upper()
    return re.sub(r'\s+', ' ', s).replace(' - ','-').replace(' -','-').replace('- ','-')

col_hs = openpyxl.utils.column_index_from_string('HS')

# ── 1. LEER LISTADO ──
listado = {}
for idx, r in enumerate(ws_l.iter_rows(min_row=3, values_only=True), start=3):
    if not any(r): continue
    raw_id = clean_str(r[3])
    if not raw_id or raw_id == 'ID LOTE': continue
    k = norm(raw_id)
    listado[k] = {
        'fila': idx, 'raw': raw_id, 'estado': clean_str(r[11]) or 'DISPONIBLE',
        'precio': to_num(r[13]), 'obs': clean_str(r[18])
    }

# ── 2. LEER VENTAS ──
ventas = {}
errores_ventas = []
for idx, r in enumerate(ws_v.iter_rows(min_row=3, values_only=True), start=3):
    if not any(r): continue
    raw_id = clean_str(r[1])
    if not raw_id or raw_id == 'ID LOTE': continue
    k = norm(raw_id)

    estado   = clean_str(r[2]) or 'DISPONIBLE'
    pv       = to_num(r[3])
    apart    = to_num(r[4])
    ci       = to_num(r[6])
    sf       = to_num(r[9])
    plazo_raw = r[10]
    plazo    = int(to_num(plazo_raw)) if to_num(plazo_raw) else 0
    cuota_raw = r[11]
    cuota    = to_num(cuota_raw)
    doc      = clean_str(r[13])
    if doc and doc.endswith('.0'): doc = doc[:-2]
    cliente  = clean_str(r[14])
    vendedor = clean_str(r[18])
    comision = to_num(r[19])
    abonos_com = to_num(r[20])
    saldo_com  = to_num(r[22])

    is_div_zero = str(cuota_raw).startswith('#DIV') or str(plazo_raw).startswith('#DIV')

    es_cedido = 'CEDIDO' in estado.upper()
    es_vendido = estado == 'VENDIDO'
    es_disp = estado in ('DISPONIBLE', 'EN NEGOCIACIÓN', 'NO APTO PARA VENTA')

    # Regla: Si vendido sin inicial ni SF -> contado
    es_contado = es_vendido and ci == 0 and sf == 0 and pv > 0
    es_financiado = es_vendido and sf > 0

    issues = []

    # A. Descuadre precio vs inicial + financiado
    if es_vendido and ci > 0 and sf > 0:
        dif_p = abs(pv - (ci + sf))
        if dif_p > 100:
            issues.append(f"Precio (${pv:,.0f}) != Inicial (${ci:,.0f}) + Financiado (${sf:,.0f}) | Dif=${dif_p:,.0f}")

    # B. Cuota mensual vs saldo financiado / plazo
    if es_financiado and plazo > 0 and cuota > 0:
        total_cuotas = cuota * plazo
        dif_c = abs(total_cuotas - sf)
        pct = (dif_c / sf * 100) if sf > 0 else 0
        if dif_c > 500000 and pct > 2:
            issues.append(f"Cuota (${cuota:,.0f} x {plazo}c = ${total_cuotas:,.0f}) vs Financiado (${sf:,.0f}) | Dif=${dif_c:,.0f}")

    # C. Vendido pero con precio 0 (sin ser cedido)
    if es_vendido and pv <= 0 and not es_cedido:
        issues.append("VENDIDO con Precio de Venta en $0")

    # D. Vendido financiado pero sin cliente o sin cédula
    if es_financiado:
        if not cliente: issues.append("Vendido financiado SIN NOMBRE de cliente")
        if not doc: issues.append("Vendido financiado SIN CÉDULA de cliente")

    # E. Error de fórmula #DIV/0! en lote que debería tener cuotas
    if is_div_zero and es_financiado:
        issues.append(f"Error de fórmula en cuota: {cuota_raw}")

    ventas[k] = {
        'fila': idx, 'raw': raw_id, 'estado': estado, 'pv': pv,
        'ci': ci, 'sf': sf, 'plazo': plazo, 'cuota': cuota,
        'doc': doc, 'cliente': cliente, 'vendedor': vendedor,
        'comision': comision, 'abonos_com': abonos_com, 'saldo_com': saldo_com,
        'es_contado': es_contado, 'es_financiado': es_financiado,
        'issues': issues
    }

# ── 3. LEER PROYECCIÓN ──
proy = {}
for idx, r in enumerate(ws_p.iter_rows(min_row=4, values_only=True), start=4):
    if not any(r): continue
    raw_id = clean_str(r[2])
    if not raw_id or raw_id == 'ID LOTE': continue
    k = norm(raw_id)
    hs_val = to_num(r[col_hs - 1]) if col_hs <= len(r) else 0.0

    # Pagos individuales
    pagos_count = 0
    pagos_suma = 0.0
    for n in range(1, 37):
        cs = 10 + (n - 1) * 6
        if cs < len(r):
            fp = r[cs+2]
            vc = to_num(r[cs+3])
            est = clean_str(r[cs+4])
            if fp and vc > 0 or (est and 'PAG' in est.upper()):
                pagos_count += 1
                pagos_suma += vc

    proy[k] = {
        'fila': idx, 'raw': raw_id, 'estado': clean_str(r[3]),
        'pv': to_num(r[4]), 'sf': to_num(r[5]), 'plazo': int(to_num(r[7])) if to_num(r[7]) else 0,
        'hs_total': hs_val, 'pagos_count': pagos_count, 'pagos_suma': pagos_suma
    }

# ── 4. VERIFICAR CASOS ESPECÍFICOS Y CRUCES ──
cruces_estado = []
cruces_precio = []
pagos_en_disponibles = []

for k, v in ventas.items():
    l = listado.get(k)
    p = proy.get(k)
    raw = v['raw']

    if l and l['estado'] != v['estado']:
        cruces_estado.append(f"{raw}: LISTADO='{l['estado']}' vs VENTAS='{v['estado']}'")

    if l and v['estado'] == 'VENDIDO' and abs(l['precio'] - v['pv']) > 1:
        cruces_precio.append({
            'lote': raw, 'cliente': v['cliente'], 'listado': l['precio'], 'ventas': v['pv'], 'dif': abs(l['precio'] - v['pv'])
        })

    if p and v['estado'] in ('DISPONIBLE', 'EN NEGOCIACIÓN') and p['hs_total'] > 0:
        pagos_en_disponibles.append({
            'lote': raw, 'estado': v['estado'], 'total_pagado': p['hs_total'], 'pagos_count': p['pagos_count']
        })

print("\n" + "="*70)
print("  REPORTE INTEGRAL DE LA HOJA DE VENTAS Y CRUCES")
print("="*70)
total_filas = len(ventas)
total_vendidos = sum(1 for v in ventas.values() if v['estado'] == 'VENDIDO')
total_disponibles = sum(1 for v in ventas.values() if v['estado'] == 'DISPONIBLE')
total_en_negociacion = sum(1 for v in ventas.values() if v['estado'] == 'EN NEGOCIACIÓN')
total_cedidos = sum(1 for v in ventas.values() if 'CEDIDO' in v['estado'].upper())
total_no_apto = sum(1 for v in ventas.values() if 'NO APTO' in v['estado'].upper())

print(f"• Total Lotes en VENTAS:              {total_filas}")
print(f"   ↳ Vendidos:                        {total_vendidos}")
print(f"   ↳ Disponibles:                     {total_disponibles}")
print(f"   ↳ En Negociación:                  {total_en_negociacion}")
print(f"   ↳ Cedidos en forma de pago:        {total_cedidos}")
print(f"   ↳ No aptos para venta:             {total_no_apto}")

ventas_con_errores = [v for v in ventas.values() if v['issues']]
print(f"\n• Inconsistencias internas en VENTAS: {len(ventas_con_errores)} casos")
if ventas_con_errores:
    for v in ventas_con_errores:
        print(f"  ❌ Fila {v['fila']}: {v['raw']} ({v['cliente']}): {v['issues']}")
else:
    print("  ✅ 100% LIMPIA: No hay descuadres de precio, cuota inicial, plazo ni cuotas.")

print(f"\n• Discrepancias de Estado (LISTADO vs VENTAS): {len(cruces_estado)} casos")
if cruces_estado:
    for ce in cruces_estado: print("  ❌", ce)
else:
    print("  ✅ 100% COINCIDENTE: Todos los estados coinciden exactamente entre hojas.")

print(f"\n• Pagos huérfanos en lotes DISPONIBLES/EN NEGOCIACIÓN: {len(pagos_en_disponibles)} casos")
for pd in pagos_en_disponibles:
    print(f"  ⚠️ {pd['lote']:<16} | Estado: {pd['estado']:<16} | Pagado en Proy: ${pd['total_pagado']:>10,.0f} ({pd['pagos_count']} cuotas)")

print(f"\n• Discrepancias de Precio Venta (LISTADO vs VENTAS): {len(cruces_precio)} casos")
for cp in cruces_precio:
    print(f"  ℹ️ {cp['lote']:<16} | Cliente: {(cp['cliente'] or '—')[:26]:<26} | LISTADO: ${cp['listado']:>11,.0f} | VENTAS: ${cp['ventas']:>11,.0f} | Dif: ${cp['dif']:>10,.0f}")

# Revisión específica de LC2-39-7
k_39_7 = norm('LC2 - 39 - 7')
p_39_7 = proy.get(k_39_7)
v_39_7 = ventas.get(k_39_7)
print(f"\n• Estado actual específico de LC2 - 39 - 7:")
print(f"   VENTAS:     Estado={v_39_7['estado'] if v_39_7 else '?'}")
print(f"   PROYECCIÓN: Estado={p_39_7['estado'] if p_39_7 else '?'}, Total Pagado HS=${p_39_7['hs_total']:,.0f if p_39_7 else 0}")
