"""
Inspeccionar LC1-6-2, LC2-39-7, LC2-44-11 y las 4 discrepancias de precio
"""
import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

FILE = 'Base de datos - La Ceiba ultima versión 2208.xlsx'
wb = openpyxl.load_workbook(FILE, data_only=True)
ws_l = wb['LISTADO']
ws_v = wb['VENTAS']
ws_p = wb['PROYECCIÓN']
ws_c = wb['CARTERA TOTAL']

# 1. 4 Discrepancias de precio LISTADO vs VENTAS
from auditoria_integral_actualizada import listado, ventas
print("=== 1. DISCREPANCIAS DE PRECIO LISTADO vs VENTAS ===")
for k, v in ventas.items():
    l = listado.get(k)
    if l and abs(l['precio'] - v['precio']) > 1 and v['estado'] == 'VENDIDO':
        print(f"• {v['raw']:<16} | Cliente: {(v['cliente'] or '—')[:26]:<26} | LISTADO: ${l['precio']:>11,.0f} | VENTAS: ${v['precio']:>11,.0f} | Dif: ${abs(l['precio']-v['precio']):>10,.0f}")

# 2. Detalles de LC1-6-2, LC2-39-7, LC2-44-11
print("\n=== 2. DETALLE DE LOTES PARTICULARES ===")
for target in ['LC1 - 6 - 2', 'LC2 - 39 - 7', 'LC2 - 44 - 11']:
    print(f"\n>>> {target} <<<")
    # VENTAS
    for r in ws_v.iter_rows(min_row=3, values_only=True):
        if r[1] and target.replace(' ','') in str(r[1]).replace(' ',''):
            print(f"  VENTAS:       Estado={r[2]} | PV={r[3]} | CI={r[6]} | SF={r[9]} | Plazo={r[10]} | Cuota={r[11]} | Cliente={r[14]} | Vendedor={r[18]}")
            break
    # PROYECCION
    for r in ws_p.iter_rows(min_row=4, values_only=True):
        if r[2] and target.replace(' ','') in str(r[2]).replace(' ',''):
            col_hs = openpyxl.utils.column_index_from_string('HS')
            print(f"  PROYECCIÓN:   Estado={r[3]} | PV={r[4]} | SF={r[5]} | Plazo={r[7]} | TOTAL PAGADO={r[col_hs-1]}")
            # Pagos
            pagos = []
            for n in range(1, 15):
                cs = 10 + (n - 1) * 6
                if cs < len(r):
                    fv, fp, vc, est = r[cs], r[cs+2], r[cs+3], r[cs+4]
                    if fv or fp or vc or est:
                        pagos.append(f"C{n}[{str(fp)[:10] if fp else 'NoPago'}: ${(vc or 0):,.0f} ({est})]")
            print(f"                Cuotas: {', '.join(pagos[:5])}")
            if len(pagos) > 5:
                print(f"                        {', '.join(pagos[5:10])}")
            break
    # CARTERA
    for r in ws_c.iter_rows(min_row=3, values_only=True):
        if r[2] and target.replace(' ','') in str(r[2]).replace(' ',''):
            print(f"  CARTERA TOT:  Estado={r[3]} | Venta={r[4]} | CI={r[5]} | SF={r[6]} | CuotasPag={r[7]} | TotPag={r[8]} | Saldo={r[9]}")
            break
