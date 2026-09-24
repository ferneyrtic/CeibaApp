"""
Inspección específica de los casos actualizados a las 15:04 (04-Sep-2026)
"""
import openpyxl, sys, re
sys.stdout.reconfigure(encoding='utf-8')

FILE = 'Base de datos - La Ceiba ultima versión 2208.xlsx'
wb = openpyxl.load_workbook(FILE, data_only=True)
ws_l = wb['LISTADO']
ws_v = wb['VENTAS']
ws_p = wb['PROYECCIÓN']
col_hs = openpyxl.utils.column_index_from_string('HS')

def norm(val):
    if not val: return ''
    return re.sub(r'\s+', ' ', str(val)).strip().upper().replace(' - ','-').replace(' -','-').replace('- ','-')

targets = [
    'LC2 - 43 - 4 Y 5', 'LC2 - 43 - 4', 'LC2 - 43 - 5',
    'LC1 - 4 - 5', 'LC1 - 16 - 8', 'LC2 - 38 - 1',
    'LC2 - 27 - 12', 'LC2 - 29 - 10', 'LC2 - 39 - 7'
]

print("="*75)
print("  REVISIÓN DE CASOS ESPECÍFICOS ACTUALIZADOS (15:04)")
print("="*75)

# Buscar en LISTADO
print("\n--- EN HOJA LISTADO ---")
for r in ws_l.iter_rows(min_row=3, values_only=True):
    raw = r[3]
    if not raw: continue
    for t in targets:
        if norm(t) == norm(raw) or norm(raw) in norm(t):
            print(f"• Lote: {raw:<16} | Estado: {r[11]} | Precio Lote(Col K): {r[10]} | Precio Venta(Col N): {r[13]} | Obs: {r[18]}")

# Buscar en VENTAS
print("\n--- EN HOJA VENTAS ---")
for r in ws_v.iter_rows(min_row=3, values_only=True):
    raw = r[1]
    if not raw: continue
    for t in targets:
        if norm(t) == norm(raw) or norm(raw) in norm(t):
            print(f"• Lote: {raw:<16} | Estado: {r[2]} | PV: {r[3]} | CI: {r[6]} | SF: {r[9]} | Plazo: {r[10]} | Cuota: {r[11]} | Cliente: {r[14]} | Doc: {r[13]}")

# Buscar en PROYECCIÓN
print("\n--- EN HOJA PROYECCIÓN ---")
for r in ws_p.iter_rows(min_row=4, values_only=True):
    raw = r[2]
    if not raw: continue
    for t in targets:
        if norm(t) == norm(raw) or norm(raw) in norm(t):
            hs = r[col_hs - 1]
            pagos = []
            for n in range(1, 10):
                cs = 10 + (n-1)*6
                fp = r[cs+2]
                vc = r[cs+3]
                est = r[cs+4]
                if fp or vc or est:
                    pagos.append(f"C{n}[Pago:{str(fp)[:10]}, Val:{vc}, Est:{est}]")
            print(f"• Lote: {raw:<16} | Estado: {r[3]} | SF: {r[5]} | Plazo: {r[7]} | Cuota: {r[9]} | Total HS: {hs}")
            if pagos:
                print(f"   ↳ Pagos:", pagos[:4])
