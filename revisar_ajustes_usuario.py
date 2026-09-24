"""
Análisis exhaustivo del archivo actualizado por el usuario (04-sep-2026)
"""
import openpyxl, sys, re, calendar
from datetime import datetime, date

sys.stdout.reconfigure(encoding='utf-8')

FILE = 'Base de datos - La Ceiba ultima versión 2208.xlsx'
print(f"Cargando archivo: {FILE}...")
wb_data = openpyxl.load_workbook(FILE, data_only=True)
wb_formula = openpyxl.load_workbook(FILE, data_only=False)

ws_l = wb_data['LISTADO']
ws_v = wb_data['VENTAS']
p_name = [s for s in wb_data.sheetnames if 'PROY' in s.upper()][0]
ws_p = wb_data[p_name]
ws_cob = wb_data['COBRANZA']
ws_cart = wb_data['CARTERA TOTAL']

def to_num(val):
    if val is None: return 0.0
    try:
        if isinstance(val, (int, float)): return float(val)
        s = str(val).replace('$','').replace(',','').strip()
        if s.startswith('#') or s in ('','None','nan'): return 0.0
        return float(s)
    except: return 0.0

def clean_str(val):
    if val is None: return None
    s = str(val).strip()
    return s if s and s.lower() not in ('none','nan','') else None

def norm(val):
    if not val: return ''
    s = str(val).strip().upper()
    return re.sub(r'\s+', ' ', s).replace(' - ','-').replace(' -','-').replace('- ','-')

col_hs = openpyxl.utils.column_index_from_string('HS')

# 1. Inspeccionar específicamente los casos mencionados por el usuario
targets = ['LC2 - 28 - 2', 'LC2 - 28 - 3', 'LC1 - 20 - 3', 'LC2 - 25 - 4', 'LC1 - 13 - 1', 'LC2 - 22 - 1']

print("\n" + "="*70)
print("  REVISIÓN ESPECÍFICA DE LOS CASOS DISCUTIDOS")
print("="*70)

for target in targets:
    t_norm = norm(target)
    print(f"\n>>> LOTE: {target} <<<")

    # En LISTADO
    l_row = None
    for r in ws_l.iter_rows(min_row=3, values_only=True):
        if r[3] and norm(r[3]) == t_norm:
            l_row = r
            break
    if l_row:
        print(f"  [LISTADO]       Estado={l_row[11]} | Precio Venta={l_row[13]} | Obs={l_row[18]}")
    else:
        print("  [LISTADO]       No encontrado")

    # En VENTAS
    v_row = None
    for r in ws_v.iter_rows(min_row=3, values_only=True):
        if r[1] and norm(r[1]) == t_norm:
            v_row = r
            break
    if v_row:
        print(f"  [VENTAS]        Estado={v_row[2]} | Precio={v_row[3]} | Inicial={v_row[6]} | SF={v_row[9]} | Plazo={v_row[10]} | Cuota={v_row[11]} | Cliente={v_row[14]} | Vendedor={v_row[18]} | Comision={v_row[19]} | Abonos={v_row[20]} | Descuento={v_row[21]}")
    else:
        print("  [VENTAS]        No encontrado")

    # En PROYECCIÓN
    p_row = None
    for r in ws_p.iter_rows(min_row=4, values_only=True):
        if r[2] and norm(r[2]) == t_norm:
            p_row = r
            break
    if p_row:
        hs_val = to_num(p_row[col_hs - 1]) if col_hs <= len(p_row) else 0.0
        print(f"  [PROYECCIÓN]    Estado={p_row[3]} | PV={p_row[4]} | SF={p_row[5]} | Plazo={p_row[7]} | TOTAL PAGADO (HS)={hs_val}")
        # Muestra cuotas con datos
        cuotas_p = []
        for n in range(1, 37):
            c_start = 10 + (n - 1) * 6
            if c_start < len(p_row):
                fv = p_row[c_start]
                fp = p_row[c_start+2]
                vc = to_num(p_row[c_start+3])
                est = p_row[c_start+4]
                if fv or fp or vc > 0 or est:
                    cuotas_p.append(f"C{n}[Venc:{str(fv)[:10]}, Pago:{str(fp)[:10]}, Val:{vc}, Est:{est}]")
        if cuotas_p:
            print(f"                  Cuotas registradas ({len(cuotas_p)}):")
            for cp in cuotas_p[:6]:
                print(f"                    ↳ {cp}")
            if len(cuotas_p) > 6:
                print(f"                    ↳ ... y {len(cuotas_p)-6} cuotas más")
        else:
            print("                  Sin cuotas registradas.")
    else:
        print("  [PROYECCIÓN]    No encontrado")

    # En CARTERA TOTAL
    c_row = None
    for r in ws_cart.iter_rows(min_row=3, values_only=True):
        if r[2] and norm(r[2]) == t_norm:
            c_row = r
            break
    if c_row:
        print(f"  [CARTERA TOTAL] Estado={c_row[3]} | Venta={c_row[4]} | Inicial={c_row[5]} | SF={c_row[6]} | Cuotas Pagadas={c_row[7]} | Total Pagado={c_row[8]} | Saldo={c_row[9]}")
    else:
        print("  [CARTERA TOTAL] No encontrado")
