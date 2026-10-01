import openpyxl, urllib.request, json, ssl, re

KEY = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InFhdGZveG1hcmRkc29jbnd5Y2d5Iiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc4MzE1NTI0MiwiZXhwIjoyMDk4NzMxMjQyfQ.cR2f5MTxRt3Tl0ZzDQJ0fFtFFnVelWq0sE5bV4chp0s'
BASE = 'https://qatfoxmarddsocnwycgy.supabase.co/rest/v1/'
ctx = ssl.create_default_context()

def get(path):
    req = urllib.request.Request(BASE + path, headers={'apikey': KEY, 'Authorization': 'Bearer ' + KEY})
    with urllib.request.urlopen(req, context=ctx) as r:
        return json.loads(r.read())

print("Consultando datos actuales en Supabase...")
db_ventas = get('ventas?select=id,lote_id,precio_venta,fecha_venta,valor_cuota_inicial,fecha_pago_cuota_inicial,abonos,saldo,estado,vendedor_nombre,lotes(id_lote),clientes(nombre,doc_cliente)&limit=1000')
db_lotes = get('lotes?select=id,id_lote,estado,precio_venta,observacion&limit=1000')

def norm(val):
    if not val: return ''
    s = str(val).strip().upper()
    return re.sub(r'\s+', ' ', s).replace(' - ','-').replace(' -','-').replace('- ','-')

db_ventas_by_lote = {norm(v['lotes']['id_lote']): v for v in db_ventas if v.get('lotes')}
db_lotes_by_id = {norm(l['id_lote']): l for l in db_lotes}

print("Cargando libro Excel...")
wb = openpyxl.load_workbook('Base de datos - La Ceiba ultima versión 2208.xlsx', data_only=True)
ws_v = wb['VENTAS']
ws_l = wb['LISTADO']
ws_p = wb[[s for s in wb.sheetnames if 'PROY' in s.upper()][0]]

print('\n=== 1. DIFERENCIAS EN LISTADO (LOTES) ===')
lotes_diffs = []
for r in ws_l.iter_rows(min_row=3, values_only=True):
    raw = r[3]
    if not raw or str(raw).strip() == 'ID LOTE': continue
    k = norm(raw)
    db_l = db_lotes_by_id.get(k)
    if not db_l:
        lotes_diffs.append((k, 'NUEVO LOTE EN EXCEL'))
        continue
    est_excel = str(r[11] or 'DISPONIBLE').strip()
    if 'CEDIDO' in est_excel.upper(): est_excel = 'VENDIDO'
    if db_l['estado'] != est_excel:
        lotes_diffs.append((k, f"Estado: DB={db_l['estado']} vs Excel={est_excel}"))

print(f"Lotes con diferencias de estado: {len(lotes_diffs)}")
for d in lotes_diffs[:10]:
    print('  ', d)

print('\n=== 2. DIFERENCIAS EN VENTAS ===')
ventas_diffs = []
for r in ws_v.iter_rows(min_row=3, values_only=True):
    raw = r[1]
    if not raw or str(raw).strip() == 'ID LOTE': continue
    k = norm(raw)
    db_v = db_ventas_by_lote.get(k)
    if not db_v:
        ventas_diffs.append((k, 'VENTA NO EN DB'))
        continue
    
    excel_estado = str(r[2] or 'DISPONIBLE').strip()
    excel_cliente = str(r[14] or '').strip()
    excel_doc = str(r[13] or '').strip()
    if excel_doc.endswith('.0'): excel_doc = excel_doc[:-2]
    excel_vendedor = str(r[18] or '').strip()
    excel_abonos = float(r[20] or 0)
    db_abonos = float(db_v.get('abonos') or 0)

    diffs = []
    if db_v['estado'] != excel_estado:
        diffs.append(f"Estado: {db_v['estado']} -> {excel_estado}")
    db_cli_nom = (db_v.get('clientes') or {}).get('nombre') or ''
    if excel_cliente and norm(db_cli_nom) != norm(excel_cliente):
        diffs.append(f"Cliente: {db_cli_nom} -> {excel_cliente}")
    if abs(db_abonos - excel_abonos) > 1:
        diffs.append(f"Abonos: {db_abonos:,.0f} -> {excel_abonos:,.0f}")
    
    if diffs:
        ventas_diffs.append((k, diffs))

print(f"Ventas con diferencias: {len(ventas_diffs)}")
for d in ventas_diffs:
    msg = ', '.join(d[1]) if isinstance(d[1], list) else str(d[1])
    print('  ', d[0], ':', msg)

print('\n=== 3. CUOTAS PAGADAS RECIENTES EN PROYECCION ===')
# Let's inspect column headers in PROYECCION
headers = [c for c in next(ws_p.iter_rows(min_row=2, max_row=2, values_only=True))]
print('Total headers in PROYECCION:', len(headers))
# Look for recent payment dates in cuotas (e.g. Sep 2026)
recent_payments = []
for r in ws_p.iter_rows(min_row=3, values_only=True):
    raw = r[1]
    if not raw or str(raw).strip() == 'ID LOTE': continue
    k = norm(raw)
    # Cuotas columns start around column 12
    # In each block of 4 cols: valor, fecha_venc, fecha_pago, estado
    for c_idx in range(12, len(r)-3, 4):
        fp = r[c_idx+2]
        est = str(r[c_idx+3] or '')
        val = r[c_idx]
        if fp:
            fp_str = str(fp)
            if '2026-09' in fp_str or '09/2026' in fp_str or '2026-08' in fp_str or '08/2026' in fp_str:
                recent_payments.append((k, (c_idx-12)//4 + 1, val, fp_str, est))

print(f"Pagos recientes detectados en PROYECCION: {len(recent_payments)}")
for p in recent_payments[-15:]:
    print('  ', p)
