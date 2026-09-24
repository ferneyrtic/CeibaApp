"""
Inspección quirúrgica de 4 puntos específicos solicitados
"""
import openpyxl, sys, re
sys.stdout.reconfigure(encoding='utf-8')

FILE = 'Base de datos - La Ceiba ultima versión 2208.xlsx'
wb = openpyxl.load_workbook(FILE, data_only=True)
wb_f = openpyxl.load_workbook(FILE, data_only=False)
ws_v = wb['VENTAS']
ws_vf = wb_f['VENTAS']
ws_l = wb['LISTADO']
ws_lf = wb_f['LISTADO']
p_name = [s for s in wb.sheetnames if 'PROY' in s.upper()][0]
ws_p = wb[p_name]
ws_pf = wb_f[p_name]
ws_cart = wb['CARTERA TOTAL']
ws_cartf = wb_f['CARTERA TOTAL']

def clean(val):
    if val is None: return None
    s = str(val).strip()
    return s if s and s.lower() not in ('none','nan','') else None

col_hs = openpyxl.utils.column_index_from_string('HS')

# ──────────────────────────────────────────────
# PUNTO 1: Rubén Fabián Bustos — revisar estado
# ──────────────────────────────────────────────
print("═"*72)
print("PUNTO 1: LOTES DE RUBÉN FABIÁN BUSTOS — Estado actualizado")
print("═"*72)
rubens = []
for idx, r in enumerate(ws_v.iter_rows(min_row=3, values_only=True), 3):
    cli = clean(r[14])
    doc = clean(r[13])
    estado = clean(r[2])
    if cli and 'RUBEN' in cli.upper() and 'BUSTOS' in cli.upper():
        rubens.append((idx, r[1], estado, cli, doc, r[9], r[10], r[11]))
    # Buscar también los sin nombre que se corrigieron
    elif not cli and clean(r[1]):
        lote = clean(r[1])
        if lote in ['LC1 - 7 - 4 ', 'LC1 - 15 - 21 ']:
            rubens.append((idx, r[1], estado, cli, doc, r[9], r[10], r[11]))

for f, lote, estado, cli, doc, sf, plazo, cuota in rubens:
    cedula_req = 'NO REQUERIDA (CEDIDO)' if estado and 'CEDIDO' in estado.upper() else ('✅ ' + str(doc) if doc else '❌ FALTA')
    print(f"Fila {f:<3} | {str(lote):<18} | Estado: {str(estado):<30} | Cliente: {str(cli)[:25] if cli else '—SIN NOMBRE—'} | Cédula: {cedula_req}")

# ──────────────────────────────────────────────
# PUNTO 2: LC1 - 6 - 2 — Celda exacta del saldo negativo
# ──────────────────────────────────────────────
print("\n" + "═"*72)
print("PUNTO 2: LC1 - 6 - 2 — Exactamente de dónde viene el saldo negativo")
print("═"*72)

# Encontrar en CARTERA TOTAL
for idx, r in enumerate(ws_cartf.iter_rows(min_row=3), 3):
    raw = r[2].value
    if raw and '6 - 2' in str(raw) and '6 - 2 ' in str(raw):
        print(f"\nHoja CARTERA TOTAL, Fila {idx}:")
        for col_idx in range(1, 11):
            cl = openpyxl.utils.get_column_letter(col_idx)
            header = ws_cartf.cell(2, col_idx).value
            formula = ws_cartf.cell(idx, col_idx).value
            value   = ws_cart.cell(idx, col_idx).value
            print(f"  Celda {cl}{idx} [{header}]:  Fórmula = {formula}  →  Valor = {value}")
        break

# Encontrar en PROYECCIÓN  
for idx, r in enumerate(ws_pf.iter_rows(min_row=4), 4):
    raw = r[2].value
    if raw and '6 - 2' in str(raw) and 'LC1' in str(raw):
        print(f"\nHoja PROYECCIÓN, Fila {idx} — columna HS (Total Pagado):")
        hs_f = ws_pf.cell(idx, col_hs).value
        hs_v = ws_p.cell(idx, col_hs).value
        print(f"  Celda HS{idx}: Fórmula = {hs_f}  →  Valor = {hs_v}")
        print(f"\n  Cuotas registradas en PROYECCIÓN (fila {idx}):")
        for n in range(1, 14):
            cs = 10 + (n-1)*6
            col_l = openpyxl.utils.get_column_letter(cs+3)
            fp = ws_p.cell(idx, cs+2).value
            vc = ws_p.cell(idx, cs+3).value
            est = ws_p.cell(idx, cs+4).value
            if fp or vc or est:
                print(f"    Cuota {n:>2} | Fecha Pago: {str(fp)[:10] if fp else '—':12} | Celda {col_l}{idx}: ${float(vc or 0):>12,.0f} | Estado: {est}")
        break

# ──────────────────────────────────────────────
# PUNTO 3: Fechas en formato incorrecto — ubicación exacta
# ──────────────────────────────────────────────
print("\n" + "═"*72)
print("PUNTO 3: FECHAS EN FORMATO INCORRECTO — Celda exacta en PROYECCIÓN")
print("═"*72)

meses_texto = {'ENERO','FEBRERO','MARZO','ABRIL','MAYO','JUNIO','JULIO','AGOSTO',
               'SEPTIEMBRE','OCTUBRE','NOVIEMBRE','DICIEMBRE'}
pattern_fecha_mal = re.compile(r'^\d{1,2}[-/]\d{1,2}[-/]\d{2,4}$|^\d{4}[-/]\d{1,2}[-/]\d{1,2}$')

fechas_malas = []
for idx, r in enumerate(ws_p.iter_rows(min_row=4, values_only=True), 4):
    if not any(r): continue
    lote = r[2]
    if not lote: continue
    for n in range(1, 37):
        cs = 10 + (n-1)*6
        if cs + 2 < len(r):
            fp_raw_idx = cs + 2
            fp_val = r[fp_raw_idx]
            if fp_val is not None:
                fp_str = str(fp_val).strip()
                col_fp = openpyxl.utils.get_column_letter(fp_raw_idx + 1)
                
                # Texto de mes
                if fp_str.upper() in meses_texto:
                    est = r[cs + 4] if cs + 4 < len(r) else None
                    is_pag = est and 'PAG' in str(est).upper()
                    fechas_malas.append((lote, n, f'{col_fp}{idx}', fp_str, 'TEXTO DE MES', is_pag))

                # Formato incorrecto con guiones o barras en distinto orden
                elif isinstance(fp_val, str) and (('-' in fp_str or '/' in fp_str) and fp_str.upper() not in ('PAGA','VENCIDA','AL DÍA')):
                    est = r[cs + 4] if cs + 4 < len(r) else None
                    is_pag = est and 'PAG' in str(est).upper()
                    fechas_malas.append((lote, n, f'{col_fp}{idx}', fp_str, 'FORMATO_TEXTO', is_pag))

for lote, nc, celda, val, tipo, is_pag in fechas_malas:
    flag = '🔴 PAGA CON FECHA MAL' if is_pag else 'ℹ️  NO AFECTA SALDO'
    print(f"  {flag} | Lote: {str(lote):<22} | Cuota {nc:>2} | Celda: {celda:<8} | Valor actual: '{val}'")

# ──────────────────────────────────────────────
# PUNTO 4: Precios diferentes LISTADO vs VENTAS — celda exacta
# ──────────────────────────────────────────────
print("\n" + "═"*72)
print("PUNTO 4: DIFERENCIAS DE PRECIO — Celda exacta en LISTADO y VENTAS")
print("═"*72)

import re as _re

def norm(val):
    if not val: return ''
    s = str(val).strip().upper()
    return _re.sub(r'\s+', ' ', s).replace(' - ','-').replace(' -','-').replace('- ','-')

def to_num(val):
    if val is None: return 0.0
    try:
        if isinstance(val, (int, float)): return float(val)
        s = str(val).replace('$','').replace(',','').strip()
        if s.startswith('#') or s in ('','None','nan'): return 0.0
        return float(s)
    except: return 0.0

listado_precios = {}
for idx, r in enumerate(ws_lf.iter_rows(min_row=3), 3):
    raw = r[3].value
    if not raw: continue
    k = norm(raw)
    precio_v = ws_l.cell(idx, 14).value  # col N = col índice 14
    col_n = openpyxl.utils.get_column_letter(14)
    listado_precios[k] = {'fila': idx, 'precio': to_num(precio_v), 'celda': f'N{idx}', 'raw': raw}

for idx, r in enumerate(ws_vf.iter_rows(min_row=3), 3):
    raw = r[1].value
    if not raw: continue
    k = norm(raw)
    estado = ws_v.cell(idx, 3).value
    if estado != 'VENDIDO': continue
    pv_v = ws_v.cell(idx, 4).value
    pv_val = to_num(pv_v)
    l = listado_precios.get(k)
    if l and abs(l['precio'] - pv_val) > 1000:
        print(f"  Lote: {str(raw):<18} | Estado: {estado}")
        print(f"    LISTADO → Hoja LISTADO, Celda {l['celda']}: ${l['precio']:,.0f}")
        print(f"    VENTAS  → Hoja VENTAS,  Celda D{idx}:  ${pv_val:,.0f}")
        print(f"    Diferencia: ${abs(l['precio'] - pv_val):,.0f}")
        print()
